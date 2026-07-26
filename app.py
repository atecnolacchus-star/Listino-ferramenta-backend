"""
app.py — backend REST per l'app Listino Ferramenta.

Avvio locale:
    pip install -r requirements.txt
    python3 seed.py
    python3 app.py
Il server parte su http://localhost:5000

Tutte le risposte sono JSON. Autenticazione tramite header:
    Authorization: Bearer <token>
Il token si ottiene da POST /api/login e dura 12 ore.
"""
import os
import io
import json
import time
import sqlite3
import secrets
from functools import wraps
from datetime import datetime

from flask import Flask, request, jsonify, g

from security import hash_password, verify_password, make_token, verify_token

try:
    import openpyxl
except ImportError:
    openpyxl = None

DB_PATH = os.path.join(os.path.dirname(__file__), 'ferramenta.db')

# Chiave di firma dei token. In produzione va impostata come variabile
# d'ambiente (SECRET_KEY) e mai lasciata al valore casuale generato qui,
# perché altrimenti tutti i token scadono ad ogni riavvio del server.
SECRET_KEY = os.environ.get('SECRET_KEY') or secrets.token_hex(32)

app = Flask(__name__)


# ---------------------------------------------------------------- utilities
def get_db():
    if 'db' not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys = ON")
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop('db', None)
    if db is not None:
        db.close()


@app.after_request
def add_cors_headers(resp):
    # CORS aperto: l'app front-end può girare da un'origine diversa
    # (es. file statico su un altro dominio o file:// in sviluppo).
    resp.headers['Access-Control-Allow-Origin'] = '*'
    resp.headers['Access-Control-Allow-Headers'] = 'Content-Type, Authorization'
    resp.headers['Access-Control-Allow-Methods'] = 'GET, POST, PATCH, DELETE, OPTIONS'
    return resp


@app.route('/api/<path:_any>', methods=['OPTIONS'])
def cors_preflight(_any):
    return ('', 204)


def err(message, status=400):
    return jsonify({'error': message}), status


def auth_required(role=None):
    """Decoratore per proteggere le route. role=None -> qualsiasi utente loggato,
    role='master' -> solo l'account master."""
    def decorator(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            header = request.headers.get('Authorization', '')
            if not header.startswith('Bearer '):
                return err('Token mancante.', 401)
            token = header[len('Bearer '):]
            payload = verify_token(token, SECRET_KEY)
            if not payload:
                return err('Token non valido o scaduto.', 401)
            db = get_db()
            acc = db.execute("SELECT * FROM accounts WHERE id=?", (payload['id'],)).fetchone()
            if not acc:
                return err('Account non trovato.', 401)
            if role and acc['role'] != role:
                return err('Permessi insufficienti.', 403)
            g.user = acc
            return fn(*args, **kwargs)
        return wrapper
    return decorator


def account_public(row):
    return {'id': row['id'], 'name': row['name'], 'role': row['role'], 'username': row['username']}


# --------------------------------------------------------------------- auth
@app.route('/api/login', methods=['POST'])
def login():
    data = request.get_json(silent=True) or {}
    username = (data.get('username') or '').strip().lower()
    password = data.get('password') or ''
    if not username or not password:
        return err('Inserisci nome utente e password.')

    db = get_db()
    acc = db.execute("SELECT * FROM accounts WHERE lower(username)=?", (username,)).fetchone()
    if not acc or not verify_password(password, acc['salt'], acc['pass_hash']):
        return err('Nome utente o password errati.', 401)

    token = make_token({'id': acc['id']}, SECRET_KEY)
    return jsonify({'token': token, 'user': account_public(acc)})


@app.route('/api/me', methods=['GET'])
@auth_required()
def me():
    return jsonify(account_public(g.user))


@app.route('/api/account/credentials', methods=['PATCH'])
@auth_required()
def change_own_credentials():
    """Ogni account (master o standard) può cambiare il proprio username e/o
    password, a patto di confermare la password attuale."""
    data = request.get_json(silent=True) or {}
    current_password = data.get('currentPassword') or ''
    new_username = (data.get('newUsername') or '').strip().lower()
    new_password = data.get('newPassword') or ''

    if not verify_password(current_password, g.user['salt'], g.user['pass_hash']):
        return err('Password attuale non corretta.', 401)
    if not new_username and not new_password:
        return err('Specifica un nuovo nome utente e/o una nuova password.')

    db = get_db()
    if new_username:
        clash = db.execute("SELECT id FROM accounts WHERE lower(username)=? AND id<>?",
                            (new_username, g.user['id'])).fetchone()
        if clash:
            return err('Nome utente già in uso.')
        db.execute("UPDATE accounts SET username=? WHERE id=?", (new_username, g.user['id']))
    if new_password:
        salt, pass_hash = hash_password(new_password)
        db.execute("UPDATE accounts SET salt=?, pass_hash=? WHERE id=?", (salt, pass_hash, g.user['id']))
    db.commit()

    acc = db.execute("SELECT * FROM accounts WHERE id=?", (g.user['id'],)).fetchone()
    # Un cambio di username invalida i vecchi riferimenti, ma non serve
    # invalidare il token corrente: resta valido fino a scadenza naturale.
    return jsonify({'ok': True, 'user': account_public(acc)})


# ---------------------------------------------------- master: gestione account
@app.route('/api/accounts', methods=['GET'])
@auth_required(role='master')
def list_accounts():
    db = get_db()
    rows = db.execute("SELECT * FROM accounts ORDER BY (role='master') DESC, id").fetchall()
    return jsonify([account_public(r) for r in rows])


@app.route('/api/accounts/<account_id>', methods=['PATCH'])
@auth_required(role='master')
def master_edit_account(account_id):
    """Il master può modificare direttamente username/password/nome di
    qualunque account, senza dover conoscere la password attuale."""
    db = get_db()
    acc = db.execute("SELECT * FROM accounts WHERE id=?", (account_id,)).fetchone()
    if not acc:
        return err('Account non trovato.', 404)

    data = request.get_json(silent=True) or {}
    new_username = (data.get('newUsername') or '').strip().lower()
    new_password = data.get('newPassword') or ''
    new_name = (data.get('newName') or '').strip()

    if new_username:
        clash = db.execute("SELECT id FROM accounts WHERE lower(username)=? AND id<>?",
                            (new_username, account_id)).fetchone()
        if clash:
            return err('Nome utente già in uso.')
        db.execute("UPDATE accounts SET username=? WHERE id=?", (new_username, account_id))
    if new_password:
        salt, pass_hash = hash_password(new_password)
        db.execute("UPDATE accounts SET salt=?, pass_hash=? WHERE id=?", (salt, pass_hash, account_id))
    if new_name:
        db.execute("UPDATE accounts SET name=? WHERE id=?", (new_name, account_id))
    db.commit()

    acc = db.execute("SELECT * FROM accounts WHERE id=?", (account_id,)).fetchone()
    return jsonify({'ok': True, 'user': account_public(acc)})


# ------------------------------------------------------------------ catalog
def serialize_catalog(db):
    cats = db.execute("SELECT * FROM categories ORDER BY id").fetchall()
    out = []
    for cat in cats:
        items = db.execute("SELECT * FROM items WHERE category_id=? ORDER BY code", (cat['id'],)).fetchall()
        out.append({
            'name': cat['name'],
            'discount': cat['discount'] or '',
            'items': [{
                'code': it['code'], 'desc': it['descr'], 'price': it['price'],
                'unit': it['unit'] or 'pz', 'discount': it['discount'] or '',
                'sub': it['sub'] or '', 'classe': it['classe'] or 'Altro',
            } for it in items]
        })
    return out


@app.route('/api/catalog', methods=['GET'])
@auth_required()
def get_catalog():
    return jsonify(serialize_catalog(get_db()))


@app.route('/api/master/classes', methods=['GET'])
@auth_required(role='master')
def get_classes():
    db = get_db()
    rows = db.execute("SELECT classe, COUNT(*) as n FROM items GROUP BY classe ORDER BY classe").fetchall()
    return jsonify([{'name': r['classe'], 'count': r['n']} for r in rows])


@app.route('/api/master/price-increase', methods=['POST'])
@auth_required(role='master')
def price_increase():
    data = request.get_json(silent=True) or {}
    classes = data.get('classes') or []
    pct = data.get('pct')
    if not classes:
        return err('Seleziona almeno una classe articolo.')
    try:
        pct = float(pct)
    except (TypeError, ValueError):
        return err('Percentuale non valida.')
    if pct == 0:
        return err('La percentuale deve essere diversa da zero.')

    db = get_db()
    placeholders = ','.join('?' * len(classes))
    rows = db.execute(f"SELECT id, price FROM items WHERE classe IN ({placeholders})", classes).fetchall()
    for r in rows:
        new_price = round(r['price'] * (1 + pct / 100), 2)
        db.execute("UPDATE items SET price=? WHERE id=?", (new_price, r['id']))
    db.commit()
    return jsonify({'updated': len(rows), 'pct': pct, 'classes': classes})


@app.route('/api/master/import', methods=['POST'])
@auth_required(role='master')
def import_excel():
    if openpyxl is None:
        return err('openpyxl non installato sul server (pip install openpyxl).', 500)
    if 'file' not in request.files:
        return err('Nessun file caricato.')
    file = request.files['file']
    wb = openpyxl.load_workbook(io.BytesIO(file.read()), read_only=True, data_only=True)
    sheet = wb.active
    rows = sheet.iter_rows(values_only=True)
    header = [str(h).strip() if h else '' for h in next(rows)]

    def col(name):
        return header.index(name) if name in header else None

    idx_code = col('Articolo')
    idx_desc = col('Descr. Art.')
    idx_price = col('Prezzo 1')
    idx_unit = col('Un. mis. 1')
    idx_classe = col('Desc. classe articolo')
    idx_sub = col('SOTTO CATEGORIA')
    idx_macro = col('MACRO CATEGORIA')
    idx_sconto = col('SCONTO SUGGERITO')

    if idx_code is None:
        return err('Colonna "Articolo" non trovata nel file: controlla la struttura.')

    db = get_db()
    updated = added = 0
    for row in rows:
        code = str(row[idx_code]).strip() if row[idx_code] else ''
        if not code:
            continue
        desc = str(row[idx_desc]).strip() if idx_desc is not None and row[idx_desc] else ''
        try:
            price = round(float(row[idx_price]), 2) if idx_price is not None and row[idx_price] not in (None, '') else 0.0
        except (TypeError, ValueError):
            price = 0.0
        unit = (str(row[idx_unit]).strip().lower() if idx_unit is not None and row[idx_unit] else 'pz') or 'pz'
        classe = (str(row[idx_classe]).strip() if idx_classe is not None and row[idx_classe] else 'Altro') or 'Altro'
        sub = (str(row[idx_sub]).strip() if idx_sub is not None and row[idx_sub] else 'Altro') or 'Altro'
        macro = (str(row[idx_macro]).strip() if idx_macro is not None and row[idx_macro] else 'Altro') or 'Altro'
        sconto_raw = row[idx_sconto] if idx_sconto is not None else None
        if isinstance(sconto_raw, (int, float)):
            discount = f"{round(sconto_raw * 100)}%"
        else:
            discount = (str(sconto_raw).strip() if sconto_raw else '') or 'Non specificato'

        existing = db.execute("SELECT id, category_id FROM items WHERE code=?", (code,)).fetchone()
        if existing:
            db.execute(
                "UPDATE items SET descr=?, price=?, unit=?, discount=?, sub=?, classe=? WHERE id=?",
                (desc, price, unit, discount, sub, classe, existing['id'])
            )
            updated += 1
        else:
            cat = db.execute("SELECT id FROM categories WHERE lower(name)=?", (macro.lower(),)).fetchone()
            if not cat:
                cur = db.execute("INSERT INTO categories (name, discount) VALUES (?,?)",
                                  (macro, 'Sconto variabile per articolo'))
                cat_id = cur.lastrowid
            else:
                cat_id = cat['id']
            db.execute(
                """INSERT INTO items (category_id, code, descr, price, unit, discount, sub, classe)
                   VALUES (?,?,?,?,?,?,?,?)""",
                (cat_id, code, desc, price, unit, discount, sub, classe)
            )
            added += 1
    db.commit()
    return jsonify({'updated': updated, 'added': added})


# ------------------------------------------------------------------- quotes
VAT_RATE = 0.22


def compute_totals(items):
    """items: lista di {qty, prezzo, sconto1, sconto2, sconto3}.
    Ogni sconto è una percentuale a cascata (es. 20% poi 10% poi 5%)."""
    base_total = 0.0
    net_total = 0.0
    for it in items:
        qty = float(it.get('qty', 0))
        prezzo = float(it.get('prezzo', 0))
        base = qty * prezzo
        net = base
        for key in ('sconto1', 'sconto2', 'sconto3'):
            pct = float(it.get(key) or 0)
            net *= (1 - pct / 100)
        base_total += base
        net_total += net
    savings = base_total - net_total
    iva = net_total * VAT_RATE
    totale_con_iva = net_total + iva
    return {
        'baseTotal': round(base_total, 2),
        'netTotal': round(net_total, 2),
        'savings': round(savings, 2),
        'ivaRate': int(VAT_RATE * 100),
        'iva': round(iva, 2),
        'totaleConIva': round(totale_con_iva, 2),
    }


@app.route('/api/quotes', methods=['GET'])
@auth_required()
def list_quotes():
    db = get_db()
    if g.user['role'] == 'master':
        rows = db.execute("SELECT * FROM quotes ORDER BY created_at DESC").fetchall()
    else:
        rows = db.execute("SELECT * FROM quotes WHERE owner_id=? ORDER BY created_at DESC", (g.user['id'],)).fetchall()
    return jsonify([quote_public(r) for r in rows])


def quote_public(row):
    return {
        'id': row['id'], 'ownerId': row['owner_id'], 'ownerName': row['owner_name'],
        'client': row['client'], 'orderInfo': json.loads(row['order_info']),
        'type': row['type'], 'status': row['status'],
        'items': json.loads(row['items']), 'totals': json.loads(row['totals']),
        'createdAt': row['created_at'],
    }


@app.route('/api/quotes', methods=['POST'])
@auth_required(role='standard')
def create_quote():
    data = request.get_json(silent=True) or {}
    items = data.get('items') or []
    order_info = data.get('orderInfo') or {}
    quote_type = data.get('type')
    client = (order_info.get('ragioneSociale') or '').strip()

    if not items:
        return err('Il carrello è vuoto.')
    if quote_type not in ('Offerta', 'Preventivo'):
        return err('Tipo non valido: usa "Offerta" o "Preventivo".')
    if not client:
        return err('Inserisci la Ragione sociale del cliente.')

    totals = compute_totals(items)
    quote_id = 'q' + str(int(time.time() * 1000)) + secrets.token_hex(2)
    created_at = datetime.utcnow().isoformat()

    db = get_db()
    db.execute(
        """INSERT INTO quotes (id, owner_id, owner_name, client, order_info, type, status, items, totals, created_at)
           VALUES (?,?,?,?,?,?,?,?,?,?)""",
        (quote_id, g.user['id'], g.user['name'], client, json.dumps(order_info),
         quote_type, 'Da confermare', json.dumps(items), json.dumps(totals), created_at)
    )
    db.commit()
    row = db.execute("SELECT * FROM quotes WHERE id=?", (quote_id,)).fetchone()
    return jsonify(quote_public(row)), 201


@app.route('/api/quotes/<quote_id>', methods=['PATCH'])
@auth_required()
def update_quote(quote_id):
    db = get_db()
    row = db.execute("SELECT * FROM quotes WHERE id=?", (quote_id,)).fetchone()
    if not row:
        return err('Preventivo non trovato.', 404)
    if g.user['role'] != 'master' and row['owner_id'] != g.user['id']:
        return err('Non puoi modificare i preventivi di un altro account.', 403)

    data = request.get_json(silent=True) or {}
    fields, values = [], []

    if 'status' in data:
        status = data.get('status')
        if status not in ('Da confermare', 'Confermato', 'Annullato'):
            return err('Stato non valido.')
        fields.append('status=?'); values.append(status)

    # Modifica completa del contenuto (articoli, intestazione, tipo).
    # Permessa solo al proprietario del preventivo (o al master) e solo se
    # non è già stato confermato.
    if 'items' in data or 'orderInfo' in data or 'type' in data:
        if row['status'] == 'Confermato' and data.get('status') != 'Da confermare':
            return err('Non è possibile modificare un preventivo già confermato: riportalo prima a "Da confermare".')
        if 'items' in data:
            items = data.get('items') or []
            if not items:
                return err('Il carrello non può essere vuoto.')
            totals = compute_totals(items)
            fields.append('items=?'); values.append(json.dumps(items))
            fields.append('totals=?'); values.append(json.dumps(totals))
        if 'orderInfo' in data:
            order_info = data.get('orderInfo') or {}
            client = (order_info.get('ragioneSociale') or '').strip()
            if not client:
                return err('Inserisci la Ragione sociale del cliente.')
            fields.append('order_info=?'); values.append(json.dumps(order_info))
            fields.append('client=?'); values.append(client)
        if 'type' in data:
            quote_type = data.get('type')
            if quote_type not in ('Offerta', 'Preventivo'):
                return err('Tipo non valido.')
            fields.append('type=?'); values.append(quote_type)

    if not fields:
        return err('Nessuna modifica specificata.')

    values.append(quote_id)
    db.execute(f"UPDATE quotes SET {', '.join(fields)} WHERE id=?", values)
    db.commit()
    row = db.execute("SELECT * FROM quotes WHERE id=?", (quote_id,)).fetchone()
    return jsonify(quote_public(row))


@app.route('/api/health', methods=['GET'])
def health():
    return jsonify({'status': 'ok', 'time': datetime.utcnow().isoformat()})


# ------------------------------------------------------------------ clients
def client_public(row):
    return {
        'id': row['id'], 'ownerId': row['owner_id'],
        'orderInfo': json.loads(row['order_info']), 'createdAt': row['created_at'],
    }


@app.route('/api/clients', methods=['GET'])
@auth_required()
def list_clients():
    db = get_db()
    rows = db.execute("SELECT * FROM clients WHERE owner_id=? ORDER BY created_at DESC", (g.user['id'],)).fetchall()
    return jsonify([client_public(r) for r in rows])


@app.route('/api/clients', methods=['POST'])
@auth_required()
def create_client():
    data = request.get_json(silent=True) or {}
    order_info = data.get('orderInfo') or {}
    if not (order_info.get('ragioneSociale') or '').strip():
        return err('Inserisci almeno la Ragione sociale.')

    client_id = 'c' + str(int(time.time() * 1000)) + secrets.token_hex(2)
    created_at = datetime.utcnow().isoformat()
    db = get_db()
    db.execute(
        "INSERT INTO clients (id, owner_id, order_info, created_at) VALUES (?,?,?,?)",
        (client_id, g.user['id'], json.dumps(order_info), created_at)
    )
    db.commit()
    row = db.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    return jsonify(client_public(row)), 201


@app.route('/api/clients/<client_id>', methods=['PATCH'])
@auth_required()
def update_client(client_id):
    db = get_db()
    row = db.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    if not row or row['owner_id'] != g.user['id']:
        return err('Anagrafica non trovata.', 404)
    data = request.get_json(silent=True) or {}
    order_info = data.get('orderInfo') or {}
    if not (order_info.get('ragioneSociale') or '').strip():
        return err('Inserisci almeno la Ragione sociale.')
    db.execute("UPDATE clients SET order_info=? WHERE id=?", (json.dumps(order_info), client_id))
    db.commit()
    row = db.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    return jsonify(client_public(row))


@app.route('/api/clients/<client_id>', methods=['DELETE'])
@auth_required()
def delete_client(client_id):
    db = get_db()
    row = db.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    if not row or row['owner_id'] != g.user['id']:
        return err('Anagrafica non trovata.', 404)
    db.execute("DELETE FROM clients WHERE id=?", (client_id,))
    db.commit()
    return jsonify({'ok': True})


# Crea il database automaticamente al primo avvio, se non esiste già.
# Girare qui (fuori da __main__) fa sì che funzioni anche con gunicorn,
# che si limita a importare questo file senza eseguire il blocco __main__.
import seed
if not os.path.exists(DB_PATH):
    seed.seed()
else:
    # Applica eventuali nuove tabelle introdotte in aggiornamenti successivi
    # (es. "clients") senza toccare account/catalogo/preventivi esistenti.
    seed.ensure_schema(DB_PATH)

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    app.run(host='0.0.0.0', port=port, debug=os.environ.get('DEBUG') == '1')
